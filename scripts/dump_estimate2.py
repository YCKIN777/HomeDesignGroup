# -*- coding: utf-8 -*-
"""读取 制安预算 121-257 行（只读）"""
from openpyxl import load_workbook

PATH = r"H:\AgentProjects\HomeDesignGroup\壹城精工私人定制产品-装修项目估算书2026.xlsx"
wb = load_workbook(PATH, data_only=True)
ws = wb["制安预算"]
for row in ws.iter_rows(min_row=121, max_row=257, values_only=False):
    vals = []
    for c in row:
        if c.value is not None:
            v = str(c.value).replace("\n", "⏎")
            if len(v) > 42:
                v = v[:42] + "…"
            vals.append(f"{c.coordinate}={v}")
    if vals:
        print(" | ".join(vals))
