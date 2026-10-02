# -*- coding: utf-8 -*-
"""解析 壹城估算书 xlsx 结构（只读）"""
from openpyxl import load_workbook

PATH = r"H:\AgentProjects\HomeDesignGroup\壹城精工私人定制产品-装修项目估算书2026.xlsx"
wb = load_workbook(PATH, data_only=True)
print("sheets:", wb.sheetnames)
for ws in wb.worksheets:
    print("=" * 60)
    print(f"[sheet] {ws.title}  dims={ws.dimensions}  max_row={ws.max_row} max_col={ws.max_column}")
    print("merged:", [str(r) for r in ws.merged_cells.ranges][:20])
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 120), values_only=False):
        vals = []
        for c in row:
            if c.value is not None:
                v = str(c.value).replace("\n", "⏎")
                if len(v) > 40:
                    v = v[:40] + "…"
                vals.append(f"{c.coordinate}={v}")
        if vals:
            print(" | ".join(vals))
