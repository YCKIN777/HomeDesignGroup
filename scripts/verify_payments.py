# -*- coding: utf-8 -*-
"""verify_payments.py — 6.1 收款闭环 + 6.3 报价导出 验证"""
import json
import sqlite3
import urllib.request

BASE = "http://127.0.0.1:8000"
DB = r"H:\AgentProjects\HomeDesignGroup\db\壹城精工.db"
PASS = 0


def ok(name, cond, extra=""):
    global PASS
    if cond:
        PASS += 1
        print(f"[PASS] {name} {extra}")
    else:
        print(f"[FAIL] {name} {extra}")


def post(path, body):
    req = urllib.request.Request(BASE + path, data=json.dumps(body).encode(),
                                 headers={"Content-Type": "application/json"})
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            return json.loads(r.read().decode())
    except urllib.error.HTTPError as e:
        return {"_http_error": e.code, "_body": e.read().decode()[:300]}


def get(path):
    with urllib.request.urlopen(BASE + path, timeout=30) as r:
        return json.loads(r.read().decode())


# ---- 前置：找一条可用的报价（王女士最新非已签约报价；无则重建种子）----
conn = sqlite3.connect(DB)
conn.row_factory = sqlite3.Row
q = conn.execute(
    "SELECT id, total, status FROM quotes WHERE project_id IN "
    "(SELECT id FROM projects WHERE customer_id IN "
    " (SELECT id FROM customers WHERE phone='19900005555')) "
    "AND status != '已签约' ORDER BY id DESC LIMIT 1").fetchone()
if not q:
    print("[setup] 重建演示报价种子…")
    conn.close()
    # 简单重建：直接调用种子脚本由外部跑；这里只提示
    raise SystemExit("请先运行 python scripts/seed_workbench_demo.py 重建演示数据")
QID = q["id"]
total = q["total"]
# 幂等化：先清空该报价的收款记录（演示数据可能已有收款）
conn.execute("DELETE FROM payments WHERE quote_id=?", (QID,))
conn.execute("DELETE FROM audit_log WHERE target LIKE 'payments#%'")
conn.commit()
conn.close()

# ---- 6.1 收款登记 ----
prog0 = get(f"/api/quotes/{QID}/payments")
ok("收款进度初始 3 阶段", len(prog0["stages"]) == 3, f"stages={len(prog0['stages'])}")
ok("签约应缴=60%总额", abs(prog0["stages"][0]["due"] - total * 0.6) < 1, f"due={prog0['stages'][0]['due']}")
ok("初始已收全 0", all(s["paid"] == 0 for s in prog0["stages"]))

# 登记签约 60% 中的一半
r1 = post(f"/api/quotes/{QID}/payments",
          {"user": "admin", "stage": "签约", "amount": round(total * 0.6 / 2, 2), "method": "转账", "note": "首笔"})
ok("登记收款成功", r1.get("id") and r1.get("amount", 0) > 0, f"id={r1.get('id')}")

prog1 = get(f"/api/quotes/{QID}/payments")
paid_s1 = prog1["stages"][0]["paid"]
ok("签约阶段已收更新", abs(paid_s1 - total * 0.3) < 1, f"paid={paid_s1}")

# 超收拦截：再收超过剩余
over = prog1["stages"][0]["remain"] + 100
r2 = post(f"/api/quotes/{QID}/payments", {"user": "admin", "stage": "签约", "amount": over})
ok("超阶段剩余拦截", r2.get("_http_error") == 400 and "剩余" in r2.get("_body", ""),
   f"http={r2.get('_http_error')}")

# 非法阶段拦截
r3 = post(f"/api/quotes/{QID}/payments", {"user": "admin", "stage": "水电", "amount": 100})
ok("非法阶段拦截", r3.get("_http_error") == 400, f"http={r3.get('_http_error')}")

# 收款明细落库
conn = sqlite3.connect(DB)
conn.row_factory = sqlite3.Row
recs = conn.execute("SELECT * FROM payments WHERE quote_id=?", (QID,)).fetchall()
ok("收款明细落库", len(recs) >= 1 and recs[0]["operator"] == "admin", f"n={len(recs)}")
conn.close()

# ---- 6.3 报价导出 ----
e = post(f"/api/quotes/{QID}/export", {"user": "admin"})
ok("导出 API 成功", e.get("file") and e.get("url"), f"file={e.get('file')} items={e.get('items')}")
if e.get("file"):
    import os
    fp = os.path.join(r"H:\AgentProjects\HomeDesignGroup\exports", e["file"])
    ok("xlsx 文件已生成", os.path.isfile(fp) and os.path.getsize(fp) > 8000,
       f"size={os.path.getsize(fp) if os.path.isfile(fp) else 0}")
    from openpyxl import load_workbook
    wb = load_workbook(fp)
    ws = wb["报价单"]
    ok("报价单 sheet 名正确", ws.title == "报价单")
    has_total = any(str(ws.cell(row=i, column=6).value) == "合计（元）" for i in range(1, ws.max_row + 1))
    ok("含合计行", has_total)
    has_warranty = any("质保" in str(ws.cell(row=i, column=1).value or "") for i in range(1, ws.max_row + 1))
    ok("含质保条款", has_warranty)
    ok("明细行数>300", ws.max_row > 300, f"max_row={ws.max_row}")

# 下载路由
try:
    with urllib.request.urlopen(BASE + "/exports/" + e["file"], timeout=15) as r:
        head = r.headers.get("Content-Disposition", "")
        ok("导出文件可下载", "attachment" in head and e["file"] in head, head)
except Exception as ex:
    ok("导出文件可下载", False, str(ex))

# ---- 项目详情附加进度 ----
import re as _re
with sqlite3.connect(DB) as _c2:
    _pid = _c2.execute("SELECT id FROM projects WHERE customer_id IN "
                       "(SELECT id FROM customers WHERE phone='19900005555') "
                       "ORDER BY id DESC LIMIT 1").fetchone()[0]
d = get(f"/api/projects/{_pid}/detail")
pp = (d.get("quote") or {}).get("payment_progress") or {}
ok("项目详情含收款进度", len(pp.get("stages", [])) == 3, f"stages={len(pp.get('stages', []))}")

print(f"\n===== {PASS} 项通过 =====")
