import os
import pandas as pd

OUTPUT_FILE = r"H:\AgentProjects\HomeDesignGroup\output\cost_summary.xlsx"


def run_cost_agent():
    import re
    import tempfile
    import time
    import zipfile
    from decimal import Decimal, ROUND_HALF_UP
    from pathlib import Path

    import openpyxl
    from openpyxl.utils.cell import range_boundaries

    output_path = Path(OUTPUT_FILE)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    template_path = Path(
        r"H:\AgentProjects\HomeDesignGroup\templates"
        r"\壹城精工私人定制产品-装修项目估算书2026.xlsx"
    )
    if not template_path.is_file():
        raise FileNotFoundError(template_path)

    cover_sheet = "\u5c01\u9762"
    hard_sheet = "\u5236\u5b89\u9884\u7b97"
    custom_sheet = "\u6210\u54c1\u5b9a\u5236"
    main_sheet = "\u4e3b\u6750\u4f30\u4ef7"

    project = {
        "owner": "\u5f85\u786e\u8ba4",
        "address": "\u5f85\u786e\u8ba4",
        "area": 78.00,
        "property_type": "villa",
    }
    quantities = {
        hard_sheet: {
            7: 42.0,
            14: 26.0,
            15: 12.0,
            17: 36.0,
            19: 12.0,
            29: 22.0,
            30: 36.0,
            37: 45.0,
            38: 24.0,
            39: 22.0,
            46: 14.0,
            51: 1.0,
            52: 5.0,
            53: 5.0,
            58: 18.0,
            65: 1.0,
            66: 4.5,
            67: 4.5,
            68: 4.5,
            75: 1.0,
            76: 18.0,
            82: 1.0,
            83: 5.0,
            84: 5.0,
            85: 5.0,
            92: 1.0,
            93: 20.0,
            109: 1.0,
            110: 4.0,
            111: 4.0,
            113: 6.0,
            116: 12.0,
            122: 1.0,
            123: 14.5,
            124: 6.0,
            125: 1.0,
            128: 14.5,
            129: 3.0,
            130: 12.0,
            131: 4.0,
            159: 1.0,
            160: 11.0,
            161: 4.0,
            162: 1.0,
            164: 11.0,
            166: 10.0,
            167: 4.0,
            171: 78.0,
            180: 1.0,
            183: 9.0,
            184: 8.0,
            185: 5.0,
            186: 18.0,
            187: 5.0,
            188: 18.0,
            196: 1.0,
            198: 78.0,
            206: 3.0,
            211: 3.0,
            214: 1.0,
            223: 1.0,
        },
        custom_sheet: {
            5: 28.0,
            6: 22.0,
            9: 8.0,
            10: 6.0,
            11: 12.0,
            16: 10.0,
            18: 12.0,
            20: 5.0,
            24: 8.0,
        },
        main_sheet: {
            5: 6.0,
            14: 42.0,
            18: 36.0,
            19: 34.0,
            22: 76.0,
            23: 16.0,
            25: 18.0,
            27: 14.0,
            29: 26.0,
            32: 4.0,
            33: 2.0,
            34: 1.0,
            35: 1.0,
            36: 2.0,
            44: 1.0,
            45: 2.0,
            46: 2.0,
            51: 2.0,
            54: 1.0,
            56: 1.0,
            57: 4.0,
            61: 12.0,
            63: 6.0,
            64: 5.0,
            66: 6.0,
            68: 70.0,
        },
    }

    workbook = openpyxl.load_workbook(template_path, data_only=False)
    try:
        for sheet_name, values in quantities.items():
            worksheet = workbook[sheet_name]
            quantity_column = 6 if sheet_name == custom_sheet else 5
            for row, quantity in values.items():
                cell = worksheet.cell(row=row, column=quantity_column)
                if cell.data_type != "f":
                    cell.value = quantity

        property_type = str(project["property_type"]).lower()
        duplex_rows = {197, 221, 222}
        villa_rows = {198, 223}
        hard_worksheet = workbook[hard_sheet]
        if property_type not in {"villa", "duplex"}:
            forbidden_rows = duplex_rows | villa_rows
        elif property_type == "villa":
            forbidden_rows = duplex_rows
        else:
            forbidden_rows = villa_rows
        for row in forbidden_rows:
            hard_worksheet.cell(row=row, column=5).value = None

        for row in range(6, 254):
            if isinstance(hard_worksheet.cell(row=row, column=6).value, (int, float)):
                hard_worksheet.cell(row=row, column=7).value = f"=F{row}*E{row}"

        hard_subtotals = {
            49: (6, 48),
            63: (51, 62),
            80: (65, 79),
            97: (82, 96),
            107: (99, 106),
            120: (109, 119),
            133: (122, 132),
            145: (135, 144),
            157: (147, 156),
            169: (159, 168),
            181: (171, 180),
            193: (183, 192),
            216: (194, 215),
            228: (218, 227),
            253: (230, 252),
        }
        for row, (start, end) in hard_subtotals.items():
            hard_worksheet.cell(row=row, column=7).value = f"=SUM(G{start}:G{end})"

        hard_worksheet["G254"].value = (
            "=G49+G63+G80+G97+G107+G120+G133+G145+G157+G169"
            "+G181+G193+G216+G228+G253"
        )
        hard_worksheet["G255"].value = "=G254*0.1"
        hard_worksheet["G256"].value = "=G254+G255"

        custom_worksheet = workbook[custom_sheet]
        for row in range(5, 26):
            if isinstance(custom_worksheet.cell(row=row, column=7).value, (int, float)):
                custom_worksheet.cell(row=row, column=8).value = f"=G{row}*F{row}"
        custom_worksheet["H26"].value = "=SUM(H5:H25)"
        custom_worksheet["H27"].value = "=H26*0.1"
        custom_worksheet["H29"].value = "=H26+H27"

        main_worksheet = workbook[main_sheet]
        for row in range(5, 77):
            if isinstance(main_worksheet.cell(row=row, column=6).value, (int, float)):
                main_worksheet.cell(row=row, column=7).value = f"=F{row}*E{row}"
        main_worksheet["G77"].value = "=SUM(G5:G76)"
        main_worksheet["G78"].value = "=G77*0.05"
        main_worksheet["G79"].value = "=G77+G78"

        hard_worksheet["G2"].value = project["area"]
        cover_worksheet = workbook[cover_sheet]
        cover_worksheet["B8"] = "\u4e1a\u4e3b\u59d3\u540d\uff1a" + project["owner"]
        cover_worksheet["B9"] = "\u5de5\u7a0b\u5730\u5740\uff1a" + project["address"]
        cover_worksheet["A13"] = (
            "\u5efa\u7b51\u9762\u79ef\uff1a" + f"{project['area']:.2f}" + "\u33a1"
        )
        cover_worksheet["E11"] = f"=[1]{hard_sheet}!G256"
        cover_worksheet["G11"] = f"={custom_sheet}!H29"
        cover_worksheet["I11"] = f"={main_sheet}!G79"
        cover_worksheet["D10"] = "=E11+G11+I11"

        formula_cache = {}

        def evaluate_cell(sheet_name, coordinate):
            key = (sheet_name, coordinate)
            if key in formula_cache:
                return formula_cache[key]

            cell = workbook[sheet_name][coordinate]
            value = cell.value
            if isinstance(value, (int, float)):
                result = float(value)
                formula_cache[key] = result
                return result
            if not isinstance(value, str) or not value.startswith("="):
                formula_cache[key] = 0.0
                return 0.0

            expression = value[1:]
            if expression.startswith("_xlfn.DISPIMG"):
                formula_cache[key] = 0.0
                return 0.0

            def replace_sum(match):
                bounds = range_boundaries(f"{match.group(1)}:{match.group(2)}")
                min_col, min_row, max_col, max_row = bounds
                total = 0.0
                for row in range(min_row, max_row + 1):
                    for column in range(min_col, max_col + 1):
                        ref_cell = workbook[sheet_name].cell(row=row, column=column)
                        total += evaluate_cell(sheet_name, ref_cell.coordinate)
                return str(total)

            expression = re.sub(
                r"SUM\(([A-Z]+\d+):([A-Z]+\d+)\)", replace_sum, expression
            ).replace("[1]", "")

            def replace_cross_sheet(match):
                return str(evaluate_cell(match.group(1), match.group(2) + match.group(3)))

            expression = re.sub(
                r"([^!+\-*/(),\s]+)!\$?([A-Z]+)\$?(\d+)",
                replace_cross_sheet,
                expression,
            )

            def replace_local(match):
                return str(evaluate_cell(sheet_name, match.group(1) + match.group(2)))

            expression = re.sub(r"\$?([A-Z]+)\$?(\d+)", replace_local, expression)
            if not re.fullmatch(r"[0-9+\-*/(). ]+", expression):
                formula_cache[key] = 0.0
                return 0.0
            result = float(eval(expression, {"__builtins__": {}}, {}))
            formula_cache[key] = result
            return result

        def rmb_upper(amount):
            cents = int(
                (Decimal(str(amount)) * 100).quantize(
                    Decimal("1"), rounding=ROUND_HALF_UP
                )
            )
            if cents == 0:
                return "\u4eba\u6c11\u5e01\u96f6\u5143\u6574"

            digits = "\u96f6\u58f9\u8d30\u53c1\u8086\u4f0d\u9646\u67d2\u634c\u7396"
            units = [
                "",
                "\u62fe",
                "\u4f70",
                "\u4edf",
                "\u4e07",
                "\u62fe\u4e07",
                "\u4f70\u4e07",
                "\u4edf\u4e07",
                "\u4ebf",
                "\u62fe\u4ebf",
                "\u4f70\u4ebf",
                "\u4edf\u4ebf",
                "\u4e07\u4ebf",
            ]
            integer, fraction = divmod(cents, 100)
            if integer:
                groups = []
                remaining = integer
                while remaining:
                    groups.append(remaining % 10000)
                    remaining //= 10000

                def group_text(group):
                    text = ""
                    zero_pending = False
                    for position in range(4):
                        digit = group % 10
                        if digit:
                            if zero_pending and text:
                                text = "\u96f6" + text
                            text = digits[digit] + units[position] + text
                            zero_pending = False
                        elif text:
                            zero_pending = True
                        group //= 10
                    return text

                big_units = [
                    "",
                    "\u4e07",
                    "\u4ebf",
                    "\u4e07\u4ebf",
                ]
                parts = []
                zero_between = False
                for index in range(len(groups) - 1, -1, -1):
                    group = groups[index]
                    if group == 0:
                        if parts:
                            zero_between = True
                        continue
                    current = group_text(group)
                    if parts and (zero_between or group < 1000):
                        parts.append("\u96f6")
                    parts.append(current + big_units[index])
                    zero_between = False
                amount_text = "".join(parts) + "\u5143"
            else:
                amount_text = "\u96f6\u5143"

            if fraction == 0:
                return "\u4eba\u6c11\u5e01" + amount_text + "\u6574"

            jiao, fen = divmod(fraction, 10)
            fraction_text = ""
            if jiao:
                fraction_text += digits[jiao] + "\u89d2"
            elif fen:
                fraction_text += "\u96f6"
            if fen:
                fraction_text += digits[fen] + "\u5206"
            return "\u4eba\u6c11\u5e01" + amount_text + fraction_text

        hard_total = evaluate_cell(hard_sheet, "G256")
        custom_total = evaluate_cell(custom_sheet, "H29")
        main_total = evaluate_cell(main_sheet, "G79")
        project_total = hard_total + custom_total + main_total
        cover_worksheet["D12"] = rmb_upper(project_total)
        cover_sheet_path = (
            f"xl/worksheets/sheet{workbook.sheetnames.index(cover_sheet) + 1}.xml"
        )
        cover_cached_values = {
            "D10": project_total,
            "E11": hard_total,
            "G11": custom_total,
            "I11": main_total,
        }
        workbook.calculation.fullCalcOnLoad = True

        rendered_fd, rendered_name = tempfile.mkstemp(
            prefix="cost_summary_", suffix=".xlsx"
        )
        os.close(rendered_fd)
        rendered_path = Path(rendered_name)
        try:
            workbook.save(rendered_path)
        finally:
            workbook.close()

        def restore_template_cellimage(source_path, rendered, destination):
            def add_content_types(xml):
                additions = [
                    b'<Default Extension="png" ContentType="image/png"/>',
                    b'<Default Extension="jpeg" ContentType="image/jpeg"/>',
                    b'<Default Extension="JPG" ContentType="image/.jpg"/>',
                    b'<Override PartName="/xl/cellimages.xml" '
                    b'ContentType="application/vnd.wps-officedocument.cellimage+xml"/>',
                ]
                for addition in additions:
                    if addition not in xml:
                        xml = xml.replace(b"</Types>", addition + b"</Types>")
                return xml

            def add_cellimage_relationship(xml):
                if b"cellimages.xml" in xml:
                    return xml
                existing_ids = [
                    int(match)
                    for match in re.findall(rb'Id="rId(\d+)"', xml)
                    if match.isdigit()
                ]
                next_id = max(existing_ids or [0]) + 1
                relationship = (
                    f'<Relationship Id="rId{next_id}" '
                    'Type="http://www.wps.cn/officeDocument/2020/cellImage" '
                    'Target="cellimages.xml"/>'
                ).encode()
                return xml.replace(
                    b"</Relationships>", relationship + b"</Relationships>"
                )

            with zipfile.ZipFile(source_path) as template, zipfile.ZipFile(
                rendered
            ) as generated:
                template_names = set(template.namelist())
                generated_names = set(generated.namelist())
                with zipfile.ZipFile(
                    destination, "w", zipfile.ZIP_DEFLATED
                ) as archived:
                    for item in generated.infolist():
                        data = generated.read(item.filename)
                        if item.filename == "[Content_Types].xml":
                            data = add_content_types(data)
                        elif item.filename == "xl/_rels/workbook.xml.rels":
                            data = add_cellimage_relationship(data)
                        elif item.filename == cover_sheet_path:
                            for coordinate, value in cover_cached_values.items():
                                pattern = (
                                    rb'(<c r="'
                                    + coordinate.encode()
                                    + rb'"[^>]*>\s*<f>.*?</f>)'
                                )
                                cached = (
                                    f"{value:.10f}".rstrip("0").rstrip(".") or "0"
                                ).encode()
                                data = re.sub(
                                    pattern,
                                    lambda match, cached=cached: match.group(1)
                                    + b"<v>"
                                    + cached
                                    + b"</v>",
                                    data,
                                    count=1,
                                )
                        archived.writestr(item, data)

                    for name in template_names:
                        if (
                            name
                            in {
                                "xl/cellimages.xml",
                                "xl/_rels/cellimages.xml.rels",
                            }
                            or (
                                name.startswith("xl/media/")
                                and name not in generated_names
                            )
                        ):
                            archived.writestr(name, template.read(name))

        try:
            restore_template_cellimage(template_path, rendered_path, output_path)
        finally:
            for attempt in range(10):
                try:
                    rendered_path.unlink(missing_ok=True)
                    break
                except PermissionError:
                    if attempt == 9:
                        raise
                    time.sleep(0.2)
    finally:
        if workbook:
            workbook.close()

    return {
        "cost_file": str(output_path),
        "total": round(project_total, 2),
    }


if __name__ == "__main__":
    result = run_cost_agent()
    print(result["cost_file"])
